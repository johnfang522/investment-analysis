"""
assemble_package.py TICKER [YYYYMMDD]

Merges the /single_stock_deep_research Word documents into one package:
the research note first, then the 8 component reports as Appendices A–H. Each
appendix starts on a new page with a bookmarked heading, the note ends with a
linked appendix index, and every page gets a "Page X of Y" footer.

    .venv/Scripts/python assemble_package.py NVDA 20261007

Reads  Outputs/{TICKER}/{ticker}_stock_deep_research_notes_{date}.docx and
       Outputs/{TICKER}/{1..8}_{ticker}_{skill}_analysis.docx
Writes Outputs/{TICKER}/{ticker}_stock_deep_research_{date}.docx

(The HTML version needs no merge: the research note page links to the
component pages through its appendix_index block.)
"""
import hashlib
import os
import sys
from copy import deepcopy
from datetime import date

from docx import Document
from docx.opc.packuri import PackURI
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.parts.image import ImagePart
from docx.shared import RGBColor

from doc_utils import apply_house_style

APPENDICES = [
    ("Appendix A", "Business Overview Analysis", "1_{t}_business_overview_analysis.docx"),
    ("Appendix B", "Leadership Analysis", "2_{t}_leadership_analysis.docx"),
    ("Appendix C", "Income Statement Analysis", "3_{t}_income_statement_analysis.docx"),
    ("Appendix D", "Balance Sheet Analysis", "4_{t}_balance_sheet_analysis.docx"),
    ("Appendix E", "Cash Flow Analysis", "5_{t}_cash_flow_analysis.docx"),
    ("Appendix F", "Business Potential Analysis", "6_{t}_business_potential_analysis.docx"),
    ("Appendix G", "Valuation Analysis", "7_{t}_valuation_analysis.docx"),
    ("Appendix H", "Technical Analysis", "8_{t}_technical_analysis.docx"),
]
IMAGE_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"


class Merger:
    def __init__(self, target):
        self.target = target
        self.images = {}        # sha1 -> ImagePart, so identical charts are stored once
        self.bookmark_id = 100

    def _image_rid(self, src_part):
        blob = src_part.blob
        sha1 = hashlib.sha1(blob).hexdigest()
        if sha1 not in self.images:
            partname = PackURI(f"/word/media/img_merged_{len(self.images) + 1}.{src_part.partname.ext}")
            self.images[sha1] = ImagePart(partname, src_part.content_type, blob)
        return self.target.part.relate_to(self.images[sha1], IMAGE_REL)

    def _body_append(self, elem):
        """Insert before the body's final sectPr — Word misplaces anything after it (lost page breaks)."""
        sect = self.target.element.body.find(qn("w:sectPr"))
        if sect is not None:
            sect.addprevious(elem)
        else:
            self.target.element.body.append(elem)

    def _bookmark(self, paragraph, name):
        self.bookmark_id += 1
        start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
        for el in (start, end):
            el.set(qn("w:id"), str(self.bookmark_id))
        start.set(qn("w:name"), name)
        paragraph._p.insert(1 if paragraph._p.pPr is not None else 0, start)
        paragraph._p.append(end)

    def add_index(self, appendices):
        """Linked list of appendices at the end of the research note."""
        self.target.add_heading("Appendices", level=2)
        for label, title, _ in appendices:
            p = self.target.add_paragraph(style="List Bullet")
            link = OxmlElement("w:hyperlink")
            link.set(qn("w:anchor"), _bookmark_name(label))
            link.set(qn("w:history"), "1")
            r, rpr = OxmlElement("w:r"), OxmlElement("w:rPr")
            color = OxmlElement("w:color"); color.set(qn("w:val"), "0563C1")
            u = OxmlElement("w:u"); u.set(qn("w:val"), "single")
            rpr.append(color); rpr.append(u); r.append(rpr)
            t = OxmlElement("w:t"); t.set(qn("xml:space"), "preserve"); t.text = f"{label} — {title}"
            r.append(t); link.append(r); p._p.append(link)

    def append(self, path, label, title):
        """New-page bookmarked appendix heading, then every body element of the source document."""
        src = Document(path)
        rid_map = {rel.rId: self._image_rid(rel.target_part)
                   for rel in src.part.rels.values() if "image" in rel.reltype}
        heading = self.target.add_heading(f"{label} — {title}", level=1)
        heading.paragraph_format.page_break_before = True
        heading.runs[0].font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
        self._bookmark(heading, _bookmark_name(label))
        for elem in src.element.body:
            if elem.tag.endswith("}sectPr"):
                continue
            copied = deepcopy(elem)
            if rid_map:
                for node in copied.iter():
                    for attr, val in list(node.attrib.items()):
                        if val in rid_map:
                            node.attrib[attr] = rid_map[val]
            self._body_append(copied)


def _bookmark_name(label):
    return label.lower().replace(" ", "_")


def _field(p, instr):
    for kind in ("begin", None, "end"):
        r = OxmlElement("w:r")
        if kind:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); r.append(fc)
        else:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {instr} "; r.append(it)
        p._p.append(r)


def add_page_numbers(doc):
    """'Page X of Y' centered in every section footer."""
    for section in doc.sections:
        section.different_first_page_header_footer = False
        footer = section.footer
        footer.is_linked_to_previous = False
        for para in footer.paragraphs:
            para.clear()
        p = footer.paragraphs[0]
        p.alignment = 1
        p.add_run("Page ")
        _field(p, "PAGE")
        p.add_run(" of ")
        _field(p, "NUMPAGES")


def assemble(ticker, day):
    ticker = ticker.upper()
    t, base = ticker.lower(), f"Outputs/{ticker}"
    note = f"{base}/{t}_stock_deep_research_notes_{day}.docx"
    if not os.path.exists(note):
        sys.exit(f"Missing research note: {note}")
    appendices = [(label, title, f"{base}/{name.format(t=t)}") for label, title, name in APPENDICES]
    missing = [p for _, _, p in appendices if not os.path.exists(p)]
    if missing:
        print("WARNING: skipping missing appendices: " + ", ".join(missing))
    appendices = [a for a in appendices if os.path.exists(a[2])]
    target = Document(note)
    merger = Merger(target)
    merger.add_index(appendices)
    for label, title, path in appendices:
        merger.append(path, label, title)
    add_page_numbers(target)
    apply_house_style(target)
    out = f"{base}/{t}_stock_deep_research_{day}.docx"
    target.save(out)
    print(f"Saved: {out}")
    return out


if __name__ == "__main__":
    if len(sys.argv) not in (2, 3):
        sys.exit(__doc__)
    assemble(sys.argv[1], sys.argv[2] if len(sys.argv) == 3 else date.today().strftime("%Y%m%d"))
