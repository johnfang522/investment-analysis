"""
report_renderer.py SPEC_JSON
report_renderer.py SPEC_JSON --format html|docx|both     (default: html)
report_renderer.py --index

Builds a self-contained interactive HTML page and, on request, a house-style Word
document (same base name) from one JSON content spec, so skills write content
(a few KB of JSON) instead of a fresh python-docx script every run. Also writes
the skill's machine-readable summary (the subagent return contract read by
/single_stock_deep_research) when the spec carries a `summary` object.

The HTML is offline-only: report_assets/report.css and report.js are inlined,
and charts are drawn from the `.chart.json` sidecar each chart_*.py writes next
to its PNG (hover tooltips, legend toggles, data table), falling back to the
embedded PNG. Every render also rebuilds Outputs/index.html, a searchable
library of every HTML report (`--index` rebuilds it alone).

    .venv/Scripts/python report_renderer.py Outputs/TSLA/5_tsla_cash_flow_spec.json

Spec shape (all keys optional except title / output / blocks):

{
  "ticker": "TSLA",
  "skill": "cash_flow",
  "title": "TSLA — Cash Flow",
  "subtitle": "Buy-side cash flow read · October 5, 2026 · Data as of: Q2 FY2026",
  "output": "Outputs/TSLA/5_tsla_cash_flow_analysis.docx",   # the .html goes next to it
  "formats": ["html"],                                        # default: html; --format / REPORT_FORMAT override
  "blocks": [ ... see BLOCK TYPES ... ],
  "summary": { ... see SUMMARY ... }
}

BLOCK TYPES — inline text in any string may use **bold** and *italic*; "\\n" breaks a line:
  {"type": "heading", "text": "At a Glance", "level": 1}
  {"type": "paragraph", "text": "..."}
  {"type": "bullets", "items": ["**Bottom line:** ...", "..."]}
  {"type": "table", "headers": [...], "rows": [[...], ...],
       "bold_rows": [6, 8],                    # 0-based data-row indexes
       "fills": [[row, col, "C6EFCE"], ...],   # cell fills: C6EFCE good / FFEB9C watch / FFC7CE bad
       "source": "SEC EDGAR (...)",            # small italic source line under the table
       "sortable": true}                        # HTML: click headers to sort (peer comps, screens)
  {"type": "chart", "path": "Outputs/TSLA/tsla_cash_flow_trend.png", "source": "SEC EDGAR — ..."}
                                               # HTML draws it from tsla_cash_flow_trend.chart.json
  {"type": "source", "text": "..."}            # standalone source line
  {"type": "variant_view", "rows": [[debate, consensus, our_read], ...],
       "source": "...", "edge": "...", "note": "..."}   # renders its own heading
  {"type": "read_through", "signal": "BULLISH|NEUTRAL|BEARISH", "dimension": "Cash-Flow-Quality",
       "conviction": 7, "so_what": "...", "what_flips": "..."}   # component skills; renders its own heading
  {"type": "verdict", "bias": "LONG|SHORT|PASS|AVOID", "conviction": 7,
       "rows": [["Current Price", "$X.XX"], ["Price Target (12-mo)", "$X.XX (+X%)"], ...],
       "bullets": ["**Justification:** ..."], "source": "..."}  # full-call skills; renders its own heading
  {"type": "income_trend", "cadence": "annual|quarterly", "source": "..."}
                                               # trend table built from chart_income_statement's
                                               # annual_/quarterly_trend_rows() — never re-type those numbers
  {"type": "page_break"}

SUMMARY (written to the spec's path with `_spec.json` -> `_summary.json`):
  {"as_of": "Q2 FY2026", "thesis_bias": "LONG|SHORT|PASS",
   "key_figures": [{"label": "TTM FCF", "value": "$6.2B", "source": "SEC EDGAR"}, ...],
   "red_flags": ["..."]}
  signal / conviction / so_what / what_flips are copied from the read_through block and
  variant_view / edge from the variant_view block, so they are never written twice.
  The first 3 key_figures also appear as stat tiles in the HTML header.
"""
import base64
import glob
import html
import json
import os
import re
import sys
from datetime import date, datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor

from chart_data import chart_data_path
from doc_utils import (setup_document, autofit_table, add_table_borders, set_row_font_size,
                       add_footnote, add_source_note, CHART_WIDTH)

ASSETS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "report_assets")
HEADER_FILL = "1F3864"
SIGNAL_COLORS = {"BULLISH": "007000", "LONG": "007000", "RISK-ON": "007000",
                 "BEARISH": "C00000", "SHORT": "C00000", "RISK-OFF": "C00000"}
READ_THROUGH_SCALE = ("Conviction scale (this dimension only): 9–10 = decisive support for the call · "
                      "7–8 = strong · 5–6 = mixed/neutral · 3–4 = weak · 1–2 = red flag")
VERDICT_SCALE = ("Conviction scale: 9–10 = highest-conviction book position · 7–8 = high · 5–6 = moderate/starter · "
                 "3–4 = low/watchlist · 1–2 = avoid or short candidate")
DISCLAIMER = ("This analysis was generated by AI with human instructions. It is provided for informational purposes "
              "only and does not constitute investment advice. Past performance is not indicative of future results. "
              "Always conduct your own research and consult a qualified financial advisor before making any "
              "investment decisions.")

_INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*)")
_LINK = re.compile(r"\[([^\]]+)\]\(([^)\s]+)\)")


# ============================== Word ==============================

def add_inline(paragraph, text, bold=False, color=None):
    """Add `text` to a paragraph, turning **x** into bold runs, *x* into italic runs and [title](url) into
    "title (url)" (Word documents keep the address visible)."""
    text = _LINK.sub(lambda m: f"{m.group(1)} ({m.group(2)})", str(text))
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2]); run.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1]); run.italic = True; run.bold = bold or None
        else:
            run = paragraph.add_run(part); run.bold = bold or None
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
    return paragraph


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _set_cell(cell, text, bold=False):
    cell.text = ""
    para = cell.paragraphs[0]
    for i, line in enumerate(str(text).split("\n")):
        if i:
            para.add_run().add_break()
        add_inline(para, line, bold=bold)


def source_line(doc, text):
    add_source_note(doc.add_paragraph(), text)


def render_table(doc, block):
    headers = block["headers"]
    tb = doc.add_table(rows=1, cols=len(headers))
    for i, h in enumerate(headers):
        cell = tb.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(str(h))
        run.bold = True; run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(cell, block.get("header_fill", HEADER_FILL))
    bold_rows = set(block.get("bold_rows", []))
    for ri, values in enumerate(block["rows"]):
        row = tb.add_row()
        for ci, v in enumerate(values):
            _set_cell(row.cells[ci], "" if v is None else v, bold=ri in bold_rows)
        set_row_font_size(row)
    for r, c, fill in block.get("fills", []):
        shade(tb.rows[r + 1].cells[c], fill.lstrip("#"))
    autofit_table(tb)
    add_table_borders(tb)
    if block.get("source"):
        source_line(doc, block["source"])


def render_chart(doc, block):
    doc.add_picture(block["path"], width=CHART_WIDTH)
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if block.get("source"):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_source_note(p, block["source"])


def render_bullets(doc, items, numbered=False):
    for item in items:
        add_inline(doc.add_paragraph(style="List Number" if numbered else "List Bullet"), item)


def render_variant_view(doc, block):
    doc.add_heading(block.get("heading", "Variant View — Consensus vs. Our Read"), level=1)
    render_table(doc, {"headers": ["Debate", "Consensus / Sell-Side", "Our Read"],
                       "rows": block["rows"], "source": block.get("source")})
    bullets = []
    if block.get("edge"):
        bullets.append(f"**The edge:** {block['edge']}")
    if block.get("note"):
        bullets.append(f"**Note:** {block['note']}")
    render_bullets(doc, bullets)


def render_read_through(doc, block):
    doc.add_heading("Read-Through to the Call", level=1)
    signal = block["signal"].upper()
    p = doc.add_paragraph()
    add_inline(p, f"Signal: {signal} (for the thesis) · {block['dimension']} Conviction "
                  f"{block['conviction']} / 10", bold=True, color=SIGNAL_COLORS.get(signal))
    for label, key in (("So what:", "so_what"), ("What flips it:", "what_flips")):
        add_inline(doc.add_paragraph(style="List Bullet"), f"{label} {block[key]}", bold=True)
    p = doc.add_paragraph(); run = p.add_run(block.get("scale", READ_THROUGH_SCALE)); run.italic = True


def render_verdict(doc, block):
    doc.add_heading(block.get("heading", "Verdict"), level=1)
    bias = block["bias"].upper()
    h = doc.add_heading(level=1)
    run = h.add_run(f"{bias} · Conviction {block['conviction']} / 10")
    if bias in SIGNAL_COLORS:
        run.font.color.rgb = RGBColor.from_string(SIGNAL_COLORS[bias])
    render_table(doc, {"headers": ["Verdict", "Call"], "rows": block["rows"], "source": block.get("source")})
    render_bullets(doc, block.get("bullets", []))
    p = doc.add_paragraph(); run = p.add_run(block.get("scale", VERDICT_SCALE)); run.italic = True


RENDERERS = {
    "heading": lambda doc, b: doc.add_heading(b["text"], level=b.get("level", 1)),
    "paragraph": lambda doc, b: add_inline(doc.add_paragraph(), b["text"], bool(b.get("color")), b.get("color")),
    "bullets": lambda doc, b: render_bullets(doc, b["items"], b.get("numbered", False)),
    "table": render_table,
    "chart": render_chart,
    "source": lambda doc, b: source_line(doc, b["text"]),
    "variant_view": render_variant_view,
    "read_through": render_read_through,
    "verdict": render_verdict,
    "page_break": lambda doc, b: doc.add_page_break(),
}

# Macro blocks (income_trend, metrics_snapshot, appendix_index) are expanded into the primitive blocks above
# before rendering — see MACRO_BUILDERS.


def _pct_str(x, signed=True):
    if x is None:
        return "N/A"
    if x == "n/m":
        return "n/m"
    return f"{x * 100:+.1f}%" if signed else f"{x * 100:.1f}%"


def _income_trend(block, ticker):
    """Annual or quarterly income-statement trend table from the same rows chart_income_statement.py drew,
    so the table and the chart can never disagree. Consensus rows are shaded and italic."""
    from chart_income_statement import annual_trend_rows, quarterly_trend_rows, margin, _price_points
    from doc_utils import fmt_value
    quarterly = block.get("cadence") == "quarterly"
    rows = quarterly_trend_rows(ticker) if quarterly else annual_trend_rows(ticker)
    px = _price_points(ticker, rows)
    out, fills = [], []
    for i, r in enumerate(rows):
        est = r["kind"] == "estimate"

        def money(key, tag=None):
            v = r.get(key)
            if v is None:
                return "N/A"
            m = margin(r, key) if tag else None
            return fmt_value(v) + (f" ({m * 100:.1f}%)" if m is not None else "")

        from datetime import datetime
        end = datetime.strptime(r["end"], "%Y-%m-%d").strftime("%m/%d/%Y") if r.get("end") else "—"
        cells = [r["period"], end, money("revenue"), _pct_str(r.get("rev_yoy")), money("gross_profit", "GM"),
                 money("operating_income", "OM"), money("net_income", "NM"),
                 f"${px[i]:,.2f}" if i in px else ("—" if est else "N/A")]
        if est:
            cells = [f"*{c}*" for c in cells]
            fills += [[i, c, "F2F2F2"] for c in range(len(cells))]
        out.append(cells)
    first = "Quarter Ended" if quarterly else "Period"
    return {"type": "table", "rows": out, "fills": fills,
            "headers": [first, "Period End", "Revenue", "Rev. YoY", "Gross Profit (GM)", "Operating Income (OM)", "Net Income (NM)", "Share Price (Period End)"],
            "source": block.get("source", "Actuals: SEC EDGAR (TTM = sum of last 4 reported quarters); share price: Yahoo Finance close on or before period end")}


_POSITIVE = ("Strong", "High quality", "Ideal", "Good", "Very conservative", "Very safe", "Very liquid", "Solid",
             "Healthy", "Undervalued", "Cheap", "High yield", "Sustainable")
_BORDERLINE = ("Decent", "Neutral", "Moderate", "Adequate", "Borderline", "Fair")
_WARNING = ("Watch", "Slow", "Thin", "Below threshold", "At risk", "Liquidity risk", "Overbought", "Expensive",
            "High risk", "Low", "Warning", "Unsustainable")
_INFO_ROWS = {"week52_low", "week52_high", "market_cap", "revenue"}


def _metric_value(key, v):
    from doc_utils import fmt_value
    if v is None:
        return "N/A"
    if key in ("current_price", "week52_low", "week52_high"):
        return f"${v:,.2f}"
    if key in ("market_cap", "revenue"):
        return fmt_value(v)
    if key in ("rev_growth", "gross_margin", "op_margin", "ni_margin", "roe", "fcf_margin", "dividend_yield",
               "payout_ratio"):
        return f"{v * 100:.1f}%"
    if key in ("de", "interest_cov"):
        return f"{v:.2f}x"
    if key in ("cur_ratio", "peg", "price_to_sales"):
        return f"{v:.2f}"
    return f"{v:.1f}"


def _label_fill(label):
    for words, fill in ((_WARNING, "FFC7CE"), (_POSITIVE, "C6EFCE"), (_BORDERLINE, "FFEB9C")):
        if any(label.startswith(w) or f" {w}" in f" {label}" for w in words):
            return fill
    return None


def _metrics_snapshot(block, ticker):
    """Financial Snapshot: Metric | Value | Description | Comments for every key in quick_stock_metrics.METRICS.
    Values, per-metric source labels and Value-cell colors come from compute_metrics() / _short_comment();
    the spec supplies only the analyst comment per key."""
    from quick_stock_metrics import METRICS, compute_metrics, _short_comment, color_current_price
    m, src = compute_metrics(ticker, with_sources=True)
    comments = block.get("comments", {})
    rows, fills = [], []
    for i, spec in enumerate(METRICS):
        k, v = spec["key"], m.get(spec["key"])
        label = _short_comment(k, v, m)
        desc = f"{spec['desc']}\nBenchmark: {spec['bench']}\nSource: {src.get(k, 'N/A')}"
        rows.append([spec["label"], _metric_value(k, v), desc, comments.get(k, "")])
        fill = None
        if v is not None and k not in _INFO_ROWS:
            if k == "current_price":
                pf = color_current_price(m)   # openpyxl fill; its pink (FFB6C1) maps to the house "bad" fill
                fill = {"FFB6C1": "FFC7CE"}.get(pf.fgColor.rgb[-6:], pf.fgColor.rgb[-6:]) if pf is not None else None
            else:
                fill = _label_fill(label)
        if fill:
            fills.append([i, 1, fill])
    missing = [k for k in comments if k not in {s["key"] for s in METRICS}]
    if missing:
        raise ValueError(f"metrics_snapshot comments use unknown keys {missing}; valid keys: "
                         f"{', '.join(s['key'] for s in METRICS)}")
    return {"type": "table", "headers": ["Metric", "Value", "Description", "Comments"], "rows": rows, "fills": fills,
            "source": block.get("source", "Per-metric source in the Description column (quick_stock_metrics."
                                          "compute_metrics); comments: analyst")}


def _appendix_index(block, ticker):
    """HTML hub: links to this ticker's component report pages with their signal and conviction."""
    t = ticker.lower()
    rows, toc_links = [], []
    for n, slug, title in ((1, "business_overview", "Business Overview"), (2, "leadership", "Leadership"),
                           (3, "income_statement", "Income Statement"), (4, "balance_sheet", "Balance Sheet"),
                           (5, "cash_flow", "Cash Flow"), (6, "business_potential", "Business Potential"),
                           (7, "valuation", "Valuation"), (8, "technical", "Technical Analysis")):
        page = f"{n}_{t}_{slug}_analysis.html"
        summ = f"Outputs/{ticker}/{n}_{t}_{slug}_summary.json"
        signal, conv = "—", "—"
        if os.path.exists(summ):
            with open(summ, encoding="utf-8") as f:
                s = json.load(f)
            signal, conv = s.get("signal") or "—", f"{s.get('conviction')}/10" if s.get("conviction") is not None else "—"
        link = f"[{title}]({page})" if os.path.exists(f"Outputs/{ticker}/{page}") else f"{title} (no HTML page yet)"
        rows.append([f"Appendix {chr(64 + n)}", link, signal, conv])
        if os.path.exists(f"Outputs/{ticker}/{page}"):
            toc_links.append([f"{chr(64 + n)}. {title}", page])
    return {"type": "table", "html_only": True, "toc_links": toc_links, "headers": ["Appendix", "Report", "Signal", "Conviction"],
            "rows": rows, "source": "Component reports in this folder; signal and conviction from each _summary.json"}


MACRO_BUILDERS = {"income_trend": _income_trend, "metrics_snapshot": _metrics_snapshot,
                  "appendix_index": _appendix_index}
MACROS = set(MACRO_BUILDERS)


def expand_blocks(spec):
    """Replace macro blocks with primitive blocks."""
    spec["blocks"] = [MACRO_BUILDERS[b["type"]](b, b.get("ticker") or spec["ticker"]) if b.get("type") in MACROS
                      else b for b in spec["blocks"]]
    return spec


def render_docx(spec):
    doc = Document()
    setup_document(doc)
    title = doc.add_paragraph(style="Title"); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(spec["title"]).bold = True
    if spec.get("subtitle"):
        sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub.add_run(spec["subtitle"]).italic = True
    for i, block in enumerate(spec["blocks"]):
        if block.get("html_only"):
            continue
        try:
            RENDERERS[block["type"]](doc, block)
        except KeyError as e:
            raise ValueError(f"block {i} ({block['type']}) is missing key {e}") from None
    add_footnote(doc)
    doc.save(spec["output"])
    print(f"Saved: {spec['output']}")


# ============================== HTML ==============================

FILL_CLASSES = {"C6EFCE": "fill-good", "FFEB9C": "fill-warn", "FFC7CE": "fill-bad", "F2F2F2": "fill-muted"}
_NUMERIC = re.compile(r"^[\s(+\-−–~≈<>≥≤]*\$?\d[\d,.]*\s*[BMKx×%]?")


def md_inline(text):
    """Escape, then **bold** -> <strong>, *italic* -> <em>, newline -> <br>."""
    out = []
    for part in _INLINE.split(str(text)):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            out.append(f"<strong>{html.escape(part[2:-2])}</strong>")
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            out.append(f"<em>{html.escape(part[1:-1])}</em>")
        else:
            out.append(_LINK.sub(lambda m: f'<a href="{m.group(2)}">{m.group(1)}</a>', html.escape(part)))
    return "".join(out).replace("\n", "<br>")


def _slug(text, used):
    base = re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-") or "section"
    slug, n = base, 2
    while slug in used:
        slug, n = f"{base}-{n}", n + 1
    used.add(slug)
    return slug


def _is_num(v):
    """Short, number-led cells ("$4.70B", "+5.3 pp", "1.51x") are right-aligned with tabular figures."""
    t = re.sub(r"\*", "", str(v)).strip()
    return len(t) <= 22 and bool(_NUMERIC.match(t)) and len(re.sub(r"[^A-Za-z]", "", t)) <= 4


def _blank(v):
    return str(v).strip() in ("", "—", "-", "N/A", "n/a", "n/m")


def html_table(block):
    headers, rows = block["headers"], block["rows"]
    num_cols = [c > 0 and any(_is_num(r[c]) for r in rows) and all(_is_num(r[c]) or _blank(r[c]) for r in rows)
                for c in range(len(headers))]
    fills = {(r, c): f for r, c, f in block.get("fills", [])}
    bold = set(block.get("bold_rows", []))
    sortable = bool(block.get("sortable")) and not bold
    th = "".join(f'<th scope="col"{" class=num" if num_cols[i] else ""}>{md_inline(h)}</th>'
                 for i, h in enumerate(headers))
    body = []
    for ri, r in enumerate(rows):
        cells = []
        for ci, v in enumerate(r):
            cls, style = (["num"] if num_cols[ci] else []), ""
            fill = fills.get((ri, ci))
            if fill:
                fill = fill.upper().lstrip("#")
                if fill in FILL_CLASSES:
                    cls.append(FILL_CLASSES[fill])
                else:
                    style = f' style="background:#{fill};color:#0b0b0b"'   # pastel row fills keep dark ink
            attr = f' class="{" ".join(cls)}"' if cls else ""
            cells.append(f"<td{attr}{style}>{md_inline('' if v is None else v)}</td>")
        body.append(f'<tr{" class=bold" if ri in bold else ""}>{"".join(cells)}</tr>')
    out = (f'<div class="table-wrap"><table{" class=sortable" if sortable else ""}><thead><tr>{th}</tr></thead>'
           f'<tbody>{"".join(body)}</tbody></table></div>')
    if block.get("source"):
        out += f'<p class="source">Source: {md_inline(block["source"])}</p>'
    return out


def html_chart(block):
    src = f'<p class="source">Source: {md_inline(block["source"])}</p>' if block.get("source") else ""
    sidecar = chart_data_path(block["path"])
    if os.path.exists(sidecar):
        with open(sidecar, encoding="utf-8") as f:
            data = json.load(f)
        payload = json.dumps(data, ensure_ascii=False).replace("</", "<\\/")
        return (f'<figure class="viz" data-chart><figcaption><span class="viz-title">'
                f'{html.escape(data.get("title", ""))}</span></figcaption>'
                f'<script type="application/json" class="chart-data">{payload}</script>{src}</figure>')
    if os.path.exists(block["path"]):
        with open(block["path"], "rb") as f:
            b64 = base64.b64encode(f.read()).decode()
        alt = html.escape(block.get("source") or os.path.basename(block["path"]))
        return (f'<figure class="viz"><img class="static-chart" alt="{alt}" '
                f'src="data:image/png;base64,{b64}">{src}</figure>')
    return f'<p class="source">[Chart not found: {html.escape(block["path"])}]</p>'


def html_bullets(items, numbered=False):
    tag = "ol" if numbered else "ul"
    return f'<{tag} class="bullets">' + "".join(f"<li>{md_inline(i)}</li>" for i in items) + f"</{tag}>"


def _signal_class(signal):
    s = str(signal).upper()
    if s in ("BULLISH", "LONG", "RISK-ON"):
        return "bullish"
    return "bearish" if s in ("BEARISH", "SHORT", "RISK-OFF") else "neutral"


def _badge(signal, suffix=""):
    return (f'<span class="badge {_signal_class(signal)}"><span class="dot"></span>'
            f'{html.escape(str(signal).upper())}{suffix}</span>')


def _assets():
    with open(os.path.join(ASSETS, "report.css"), encoding="utf-8") as f:
        css = f.read()
    with open(os.path.join(ASSETS, "report.js"), encoding="utf-8") as f:
        js = f.read()
    return css, js


def _page(title, body, meta=None):
    css, js = _assets()
    meta_tag = (f'<meta name="report-meta" content="{html.escape(json.dumps(meta, ensure_ascii=False))}">\n'
                if meta else "")
    return (f'<!doctype html>\n<html lang="en">\n<head>\n<meta charset="utf-8">\n'
            f'<meta name="viewport" content="width=device-width, initial-scale=1">\n{meta_tag}'
            f'<title>{html.escape(title)}</title>\n<style>{css}</style>\n</head>\n<body>\n{body}\n'
            f'<script>{js}</script>\n</body>\n</html>\n')


def render_html(spec, out_path, summary):
    used, toc, parts, toc_sub = set(), [], [], {}

    def heading(text, level=1):
        tag = "h2" if level <= 1 else "h3"
        slug = _slug(text, used)
        if tag == "h2":
            toc.append((slug, text))
        return f'<{tag} id="{slug}">{md_inline(text)}</{tag}>'

    for i, block in enumerate(spec["blocks"]):
        kind = block["type"]
        try:
            if kind == "heading":
                parts.append(heading(block["text"], block.get("level", 1)))
            elif kind == "paragraph":
                if block.get("color"):   # verdict-style line: bold, tinted by the call (status colors in CSS)
                    tone = {"007000": "good", "C00000": "bad", "BF8F00": "warn", "FF8C00": "serious"}.get(
                        block["color"].upper().lstrip("#"), "neutral")
                    parts.append(f'<p class="call-line {tone}"><strong>{md_inline(block["text"])}</strong></p>')
                else:
                    parts.append(f"<p>{md_inline(block['text'])}</p>")
            elif kind == "bullets":
                parts.append(html_bullets(block["items"], block.get("numbered", False)))
            elif kind == "table":
                parts.append(html_table(block))
                if block.get("toc_links") and toc:   # appendix hub: list each appendix under its heading in the sidebar
                    toc_sub[toc[-1][0]] = block["toc_links"]
            elif kind == "chart":
                parts.append(html_chart(block))
            elif kind == "source":
                parts.append(f'<p class="source">Source: {md_inline(block["text"])}</p>')
            elif kind == "page_break":
                parts.append('<hr class="page-break">')
            elif kind == "variant_view":
                parts.append(heading(block.get("heading", "Variant View — Consensus vs. Our Read")))
                parts.append(html_table({"headers": ["Debate", "Consensus / Sell-Side", "Our Read"],
                                         "rows": block["rows"], "source": block.get("source")}))
                if block.get("edge"):
                    parts.append(f'<div class="edge"><strong>The edge:</strong> {md_inline(block["edge"])}</div>')
                if block.get("note"):
                    parts.append(f'<p><strong>Note:</strong> {md_inline(block["note"])}</p>')
            elif kind == "read_through":
                cls = _signal_class(block["signal"])
                parts.append(heading("Read-Through to the Call"))
                parts.append(
                    f'<div class="callout {cls}"><div class="signal">{_badge(block["signal"])} '
                    f'{html.escape(block["dimension"])} conviction {html.escape(str(block["conviction"]))} / 10</div>'
                    + html_bullets([f"**So what:** {block['so_what']}", f"**What flips it:** {block['what_flips']}"])
                    + f'<div class="scale">{html.escape(block.get("scale", READ_THROUGH_SCALE))}</div></div>')
            elif kind == "verdict":
                cls = _signal_class(block["bias"])
                parts.append(heading(block.get("heading", "Verdict")))
                tiles = "".join(f'<div class="stat"><div class="label">{md_inline(k)}</div>'
                                f'<div class="value">{md_inline(v)}</div></div>' for k, v in block["rows"])
                parts.append(
                    f'<div class="callout verdict {cls}"><div class="signal">{_badge(block["bias"])} '
                    f'Conviction {html.escape(str(block["conviction"]))} / 10</div>'
                    f'<div class="call-strip">{tiles}</div>'
                    + (html_bullets(block["bullets"]) if block.get("bullets") else "")
                    + (f'<p class="source">Source: {md_inline(block["source"])}</p>' if block.get("source") else "")
                    + f'<div class="scale">{html.escape(block.get("scale", VERDICT_SCALE))}</div></div>')
        except KeyError as e:
            raise ValueError(f"block {i} ({kind}) is missing key {e}") from None

    strip = []
    if summary.get("signal"):
        strip.append(f'<div class="stat"><div class="label">Signal</div><div class="value">'
                     f'{_badge(summary["signal"])}</div></div>')
    if summary.get("conviction") is not None:
        c = summary["conviction"]
        try:
            pct = max(0.0, min(100.0, float(c) * 10))
        except (TypeError, ValueError):
            pct = 0.0
        strip.append(f'<div class="stat"><div class="label">Conviction</div><div class="value">'
                     f'{html.escape(str(c))} / 10</div><div class="meter"><span style="width:{pct:.0f}%"></span>'
                     f'</div></div>')
    if summary.get("thesis_bias"):
        strip.append(f'<div class="stat"><div class="label">Thesis bias</div><div class="value">'
                     f'{html.escape(summary["thesis_bias"])}</div></div>')
    for kf in summary.get("key_figures", [])[:3]:
        main_v, _, rest = str(kf["value"]).partition(" (")
        sub = f'<div class="sub">({html.escape(rest)}</div>' if rest else ""
        strip.append(f'<div class="stat" title="Source: {html.escape(kf.get("source", ""))}">'
                     f'<div class="label">{html.escape(kf["label"])}</div>'
                     f'<div class="value">{html.escape(main_v)}</div>{sub}</div>')

    eyebrow = " · ".join(x for x in (spec.get("ticker"), (spec.get("skill") or "").replace("_", " ").title()) if x)
    index_rel = os.path.relpath(os.path.join("Outputs", "index.html"),
                                os.path.dirname(out_path) or ".").replace(os.sep, "/")
    def toc_item(slug, t):
        sub = "".join(f'<li><a href="{html.escape(h)}">{html.escape(lbl)}</a></li>' for lbl, h in toc_sub.get(slug, []))
        return f'<li><a href="#{slug}">{html.escape(t)}</a>{f"<ol class=sub>{sub}</ol>" if sub else ""}</li>'
    toc_html = "".join(toc_item(slug, t) for slug, t in toc)
    subtitle = f'<p class="subtitle">{md_inline(spec["subtitle"])}</p>' if spec.get("subtitle") else ""
    call_strip = f'<div class="call-strip">{"".join(strip)}</div>' if strip else ""
    shell_cls = "shell wide" if spec.get("layout") == "wide" else "shell"
    back_rel = _package_note_for(out_path) or index_rel   # a component report goes back to its research package
    back = (f'<a class="back-fab" href="{back_rel}" title="Back to the previous page (drag to move)">'
            f'<span aria-hidden="true">←</span> Back</a>\n')
    body = (back + f'<div class="{shell_cls}">\n<nav class="toc" aria-label="Contents"><a class="toc-home" href="{index_rel}">'
            f'← Investment Research Library</a><ol>{toc_html}</ol></nav>\n<main>\n<header class="report-head">'
            f'<div class="eyebrow"><span>{html.escape(eyebrow)}</span>'
            f'<button type="button" class="theme-toggle">Dark mode</button></div>'
            f'<h1>{md_inline(spec["title"])}</h1>{subtitle}{call_strip}</header>\n'
            + "\n".join(parts)
            + f'\n<footer class="disclaimer">{html.escape(DISCLAIMER)}</footer>\n</main>\n</div>')
    meta = {"ticker": spec.get("ticker"), "skill": spec.get("skill"), "title": spec["title"],
            "signal": summary.get("signal"), "conviction": summary.get("conviction"),
            "as_of": summary.get("as_of"), "rendered": date.today().isoformat()}
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(_page(spec["title"], body, meta))
    print(f"Saved: {out_path}")
    if _NOTE_RE.search(out_path):
        _retarget_back_links(out_path)


def _package_note_for(out_path):
    """Newest deep-research note page in the same folder when out_path is one of its component reports, else None."""
    if not _COMPONENT_RE.search(out_path.replace(os.sep, "/")):
        return None
    notes = sorted(f for f in os.listdir(os.path.dirname(out_path) or ".") if _NOTE_RE.search(f))
    return notes[-1] if notes else None


def _retarget_back_links(note_path):
    """After a research note renders, point its component pages' Back button at it (components render first)."""
    folder, note = os.path.split(note_path)
    for f in os.listdir(folder or "."):
        if not _COMPONENT_RE.search(f):
            continue
        path = os.path.join(folder, f)
        with open(path, encoding="utf-8") as fh:
            text = fh.read()
        new = re.sub(r'(<a class="back-fab" href=")[^"]*(")', lambda m: m.group(1) + note + m.group(2), text, count=1)
        if new != text:
            with open(path, "w", encoding="utf-8") as fh:
                fh.write(new)


def _read_meta(path):
    with open(path, encoding="utf-8", errors="ignore") as f:
        head = f.read(4096)
    m = re.search(r'<meta name="report-meta" content="([^"]*)"', head)
    if m:
        try:
            return json.loads(html.unescape(m.group(1)))
        except ValueError:
            pass
    t = re.search(r"<title>(.*?)</title>", head, re.S)
    return {"title": html.unescape(t.group(1)).strip() if t else os.path.basename(path)}


NEW_DAYS = 7   # reports younger than this get a "New" tag on the library page


def _company_name(folder, root="Outputs"):
    """Full company name for a ticker folder (Yahoo longName), formatted "Tesla, Inc. (TSLA)"; the ticker alone if unknown."""
    try:
        with open(os.path.join(root, folder, f"{folder.lower()}_quick_metrics.json"), encoding="utf-8") as f:
            name = json.load(f).get("longName")
    except (OSError, ValueError):
        name = None
    return f"{name} ({folder})" if name else folder


_COMPONENT_RE = re.compile(r"(?:^|/)([1-8])_[^/]+_analysis\.html$")
_NOTE_RE = re.compile(r"_stock_deep_research_notes_(\d{8})\.html$")


def _index_item(m, children="", strip_prefix=None):
    badge = ""
    if m.get("signal"):
        conv = f' {m["conviction"]}/10' if m.get("conviction") is not None else ""
        badge = _badge(m["signal"], html.escape(conv))
    when = datetime.fromtimestamp(m["mtime"]).strftime("%Y-%m-%d %H:%M")
    title = m.get("title") or m["href"]
    if strip_prefix and title.startswith(strip_prefix):
        title = title[len(strip_prefix):]
    fresh = (datetime.now().timestamp() - m["mtime"]) < NEW_DAYS * 86400
    tag = '<span class="new-tag">New</span>' if fresh else ""
    row = (f'<div class="row"><span class="ttl"><a href="{html.escape(m["href"])}">{html.escape(title)}</a>{tag}</span>'
           f'<span class="meta">{badge} {when}</span></div>')
    return f'<li{" class=pkg" if children else ""}>{row}{children}</li>'


def _card(folder, items):
    """One ticker card: the newest deep-research package on top with its 8 component reports as a sub-list,
    then older packages and other reports (newest first). Header shows the last research run."""
    notes = sorted((m for m in items if _NOTE_RE.search(m["href"])), key=lambda m: m["href"], reverse=True)
    comps = sorted((m for m in items if _COMPONENT_RE.search(m["href"])),
                   key=lambda m: _COMPONENT_RE.search(m["href"]).group(1))
    others = sorted((m for m in items if m not in notes and m not in comps), key=lambda m: -m["mtime"])
    # top-level entries (packages, standalone reports) newest first; a package carries its 8 components (1-8)
    entries = []
    if notes:
        sub = (f'<ul class="sub">{"".join(_index_item(c, strip_prefix=f"{folder} — ") for c in comps)}</ul>'
               if comps else "")
        entries.append((notes[0]["mtime"], _index_item(notes[0], sub)))
        entries += [(m["mtime"], _index_item(m)) for m in notes[1:]]
    else:
        entries += [(m["mtime"], _index_item(m)) for m in comps]
    entries += [(m["mtime"], _index_item(m)) for m in others]
    lis = [li for _, li in sorted(entries, key=lambda e: -e[0])]
    last = max(m["mtime"] for m in (notes[:1] or items))
    label = "Last research run" if notes else "Last run"
    stamp = datetime.fromtimestamp(last).strftime("%Y-%m-%d %H:%M")
    top = (notes[:1] or [None])[0]
    badge = ""
    if top and top.get("signal"):
        conv = f' {top["conviction"]}/10' if top.get("conviction") is not None else ""
        badge = _badge(top["signal"], html.escape(conv))
    return last, (f'<section class="card"><details><summary><span class="card-head"><h2>{html.escape(_company_name(folder))}</h2></span>'
                  f'<span class="card-meta">{badge}<span class="last-run">{label}: <time datetime="'
                  f'{datetime.fromtimestamp(last).isoformat()}">{stamp}</time></span></span></summary>'
                  f'<ul>{"".join(lis)}</ul></details></section>')


STAGES = [
    (1, "Market Conditions", "Assess broad market sentiment and risk before deploying capital.",
     ("market_sentiment",)),
    (2, "Theme Discovery", "Identify the value chain for a macro trend and surface candidate stocks at each layer.",
     ("theme_discovery", "emerging_industry", "industry_trend", "industry_deep_dive", "ai_company_deep_dive",
      "company_deep_dive", "multibagger")),
    (3, "Quick Filter", "Screen candidates on financial quality before committing to deep research.",
     ("quick_stock_metrics",)),
    (4, "Individual Stock Analysis", "Deep-dive on specific names across all dimensions, culminating in a research note.",
     ()),
]


def _stage_of(meta):
    """Stage number for a report, from its skill name (stage 4 is the default for any single-stock report)."""
    key = (meta.get("skill") or "") + " " + os.path.basename(meta["href"])
    for num, _, _, keys in STAGES:
        if any(k in key for k in keys):
            return num
    return 4


def _skill_card(label, items):
    """Stage 1-3 card: one skill's reports, newest first, with the ticker in front when there is one."""
    items = sorted(items, key=lambda m: -m["mtime"])
    lis = []
    for m in items:
        title = m.get("title") or m["href"]
        if m.get("ticker") and not title.upper().startswith(m["ticker"].upper()):
            title = f'{m["ticker"]} — {title}'
        lis.append(_index_item(dict(m, title=title)))
    last = datetime.fromtimestamp(items[0]["mtime"]).strftime("%Y-%m-%d %H:%M")
    return (f'<section class="card"><h2>{html.escape(label)}</h2><p class="last-run">Last run: '
            f'<time>{last}</time></p><ul>{"".join(lis)}</ul></section>')


def _skill_label(meta):
    return (meta.get("skill") or os.path.basename(meta["href"]).rsplit(".", 1)[0]).replace("_", " ").title()


_DATED_RE = re.compile(r"^(.*)_(\d{8})\.html$")


def prune_superseded(root="Outputs"):
    """A skill re-run for the same subject on a later day replaces the earlier report: for every dated report
    (`<name>_YYYYMMDD.html`) keep only the newest date and delete the older ones with their `.docx`, `_spec.json`
    and `_summary.json` (and `.xlsx` for the metrics workbook). Charts and data JSON are left alone."""
    groups = {}
    for path in glob.glob(os.path.join(root, "**", "*.html"), recursive=True):
        m = _DATED_RE.match(path.replace(os.sep, "/"))
        if m:
            groups.setdefault(m.group(1), []).append((m.group(2), path))
    removed = []
    for prefix, found in groups.items():
        found.sort()
        for day, path in found[:-1]:
            stem = path[:-len(".html")]
            for ext in (".html", ".docx", ".xlsx", "_spec.json", "_summary.json"):
                if os.path.exists(stem + ext):
                    os.remove(stem + ext)
                    removed.append(stem + ext)
    for r in removed:
        print(f"Removed superseded: {r}")


def build_index(root="Outputs"):
    """Rebuild Outputs/index.html: a left navigation and the reports organised by the 4 workflow stages
    (market conditions, theme discovery, quick filter, then one card per ticker in A→Z order)."""
    prune_superseded(root)
    metas = []
    for path in glob.glob(os.path.join(root, "**", "*.html"), recursive=True):
        if os.path.basename(path) == "index.html":
            continue
        rel = os.path.relpath(path, root).replace(os.sep, "/")
        meta = _read_meta(path)
        meta.update(href=rel, mtime=os.path.getmtime(path), folder=rel.split("/")[0] if "/" in rel else "")
        metas.append(meta)
    by_stage = {n: [] for n, *_ in STAGES}
    for m in metas:
        by_stage[_stage_of(m)].append(m)

    sections, nav = [], []
    for num, name, desc, _ in STAGES:
        items = by_stage[num]
        sub = ""
        if num == 4:
            groups = {}
            for m in items:
                groups.setdefault(m["folder"] or "Other", []).append(m)
            order = sorted(groups, key=lambda f: (f == "Other", f.upper()))
            cards = "".join(_card(f, groups[f])[1].replace("<section class=\"card\">",
                            f'<section class="card" id="t-{_slug(f, set())}">', 1) for f in order)
            sub = "".join(f'<li><a href="#t-{_slug(f, set())}">{html.escape(_company_name(f))}</a></li>' for f in order)
        else:
            groups = {}
            for m in items:
                groups.setdefault(_skill_label(m), []).append(m)
            cards = "".join(_skill_card(k, v) for k, v in sorted(groups.items()))
        body = f'<div class="cards">{cards}</div>' if cards else '<p class="empty">No reports yet.</p>'
        sections.append(f'<section class="stage" id="stage-{num}"><h2 class="stage-title"><span class="stage-no">'
                        f'Stage {num}</span> — {html.escape(name)}</h2><p class="stage-desc">{html.escape(desc)}</p>'
                        f'{body}</section>')
        nav.append(f'<li><a href="#stage-{num}"><b>Stage {num}</b> — {html.escape(name)}</a>{f"<ol class=sub>{sub}</ol>" if sub else ""}</li>')

    body = (f'<div class="shell library"><nav class="toc" aria-label="Stages"><a class="toc-home" href="#">'
            f'Investment Research Library</a><ol>{"".join(nav)}</ol></nav><main>'
            f'<div class="library-head"><div><h1>Investment Research Library</h1>'
            f'<p class="subtitle" style="color:var(--muted);margin:4px 0 0">Rebuilt '
            f'{date.today().isoformat()}</p></div><div style="display:flex;flex-wrap:wrap;gap:10px;align-items:center">'
            f'<input id="library-search" type="search" placeholder="Filter by ticker or title" '
            f'aria-label="Filter reports"><button type="button" class="theme-toggle">Dark mode</button></div></div>'
            f'{"".join(sections)}</main></div>')
    out = os.path.join(root, "index.html")
    with open(out, "w", encoding="utf-8") as f:
        f.write(_page("Investment Research Library", body))
    print(f"Saved: {out}")


# ============================== entry point ==============================

def build_summary(spec):
    """Merge the spec's `summary` with the read-through and variant-view blocks."""
    summary = {"ticker": spec.get("ticker"), "skill": spec.get("skill"), "document": spec["output"]}
    summary.update(spec.get("summary", {}))
    for block in spec["blocks"]:
        if block["type"] == "read_through":
            summary.update({k: block[k] for k in ("signal", "conviction", "so_what", "what_flips")})
            summary["signal"] = summary["signal"].upper()
        elif block["type"] == "verdict":
            summary.update(signal=block["bias"].upper(), conviction=block["conviction"],
                           verdict={str(k).replace("*", ""): str(v).replace("*", "") for k, v in block["rows"]})
        elif block["type"] == "variant_view":
            summary["variant_view"] = [dict(zip(("debate", "consensus", "our_read"), r)) for r in block["rows"]]
            summary["edge"] = block.get("edge")
    return summary


def validate(spec):
    for i, block in enumerate(spec["blocks"]):
        kind = block.get("type")
        if kind not in RENDERERS and kind not in MACROS:
            raise ValueError(f"block {i}: unknown type {kind!r} (valid: {', '.join(list(RENDERERS) + sorted(MACROS))})")
        if kind in ("table", "variant_view", "verdict"):
            width = len(block["headers"]) if kind == "table" else 3 if kind == "variant_view" else 2
            for ri, row in enumerate(block["rows"]):
                if len(row) != width:
                    raise ValueError(f"block {i} ({kind}) row {ri} has {len(row)} cells, expected {width}: {row}")


FORMAT_CHOICES = {"html": ["html"], "docx": ["docx"], "both": ["docx", "html"]}


def render(spec_path, fmt=None):
    """fmt: 'html' (default), 'docx' or 'both'. Precedence: --format flag, then the
    REPORT_FORMAT env var, then the spec's own "formats", then html only."""
    with open(spec_path, encoding="utf-8") as f:
        spec = json.load(f)
    validate(spec)
    expand_blocks(spec)
    fmt = fmt or os.environ.get("REPORT_FORMAT")
    if fmt:
        if fmt not in FORMAT_CHOICES:
            sys.exit(f"Unknown format {fmt!r}; choose html, docx or both")
        formats = FORMAT_CHOICES[fmt]
    else:
        formats = spec.get("formats", ["html"])
    summary = build_summary(spec)
    if "docx" in formats:
        render_docx(spec)
    if "html" in formats:
        render_html(spec, os.path.splitext(spec["output"])[0] + ".html", summary)
        if os.path.isdir("Outputs"):
            build_index("Outputs")
    if "summary" in spec:
        summary_path = re.sub(r"_spec\.json$", "_summary.json", spec_path)
        if summary_path == spec_path:
            summary_path = spec_path.rsplit(".", 1)[0] + "_summary.json"
        with open(summary_path, "w", encoding="utf-8") as f:
            json.dump(summary, f, indent=2, ensure_ascii=False)
        print(f"Saved: {summary_path}")


if __name__ == "__main__":
    if len(sys.argv) == 2 and sys.argv[1] == "--index":
        build_index("Outputs")
    elif len(sys.argv) == 2:
        render(sys.argv[1])
    elif len(sys.argv) == 4 and sys.argv[2] == "--format":
        render(sys.argv[1], sys.argv[3])
    else:
        sys.exit(__doc__)
